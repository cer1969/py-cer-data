# CRISTIAN ECHEVERRÍA RABÍ
#
# Versión mejorada de 'ox'

import openpyxl

#---------------------------------------------------------------------------------------------------

class Book(object):

    def __init__(self, filepath, read_only=True, data_only=True):
        self._wb = openpyxl.load_workbook(filepath, read_only=read_only, data_only=data_only)

    def get_sheet(self, name):
        return Sheet(self, name)
    
    def get_table_data(self, sheet: str, table: str, dict_class=dict):
        """
        Returns a list of dictionaries containing the data from the specified table in 
        the specified sheet.
        The workbook must be opened with read_only=False.

        Parameters:
        sheet (str): The name of the sheet to read from.
        table (str): The name of the table to read from.
        dict_class (type): The type of dictionary to use for the result (default is dict).

        Returns:
        list of dict_class: A list of dictionaries containing the data from the specified 
        table in the specified sheet.
        """
        return Sheet(self, sheet).get_table_data(table, dict_class)


#---------------------------------------------------------------------------------------------------

class Sheet(object):

    def __init__(self, book, name):
        self._sh = book._wb[name]

    def iter_data(self, rows=(1, None), cols=(1, None)):
        return self._sh.iter_rows(min_row=rows[0], max_row=rows[1], min_col=cols[0], max_col=cols[1])

    def get_data(self, rows, cols, headers_row=None, dict_class=dict):
        """headers_row: Si headers_row es None se considera que rows incluye los headers
           dict_class: Permite cambiar la clase de diccionario (Ej: munch)
        """
        it = self.iter_data(rows, cols)
        
        if headers_row is None:
            headers = [c.value for c in next(it)]
        else:
            it_headers = self.iter_data((headers_row, headers_row), cols)
            headers = [c.value for c in next(it_headers)]

        data = []
        for r in it:
            data.append(dict_class(zip(headers, [c.value for c in r])))
        return data
    
    def get_table_data(self, table, dict_class=dict):
        # REQUIERE read_only=False
        t = self._sh.tables[table]
        data = self._sh[t.ref]

        filas = [[celda.value for celda in fila] for fila in data]
        headers = filas[0]

        return [dict_class(zip(headers, fila)) for fila in filas[1:]]